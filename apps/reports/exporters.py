"""
Export utilities for CSV, Excel (XLSX), and PDF generation.
"""

import csv
from datetime import datetime

from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer


REPORT_COLUMNS = [
    'Sr. No.',
    'Student ID',
    'Student Name',
    'Family ID',
    'Family / Firm Name',
    'Responsible Person',
    'Village / City',
    'Standard',
    'Medium',
    'Percentage / CGPA',
    'Special Achievement',
    'Award',
    'Status',
]


def get_report_rows(registrations):
    """Convert registration queryset to list of row data."""
    rows = []
    for idx, reg in enumerate(registrations, 1):
        awards = ', '.join(a.category.name for a in reg.awards.all()) if reg.awards.exists() else ''
        rows.append([
            idx,
            reg.student.student_id,
            reg.student.full_name,
            reg.student.family.family_id,
            reg.student.family.family_name,
            reg.student.family.responsible_person_name,
            reg.student.family.village_city,
            reg.standard,
            reg.get_medium_display(),
            str(reg.percentage_or_cgpa),
            reg.special_achievement or '',
            awards,
            reg.get_status_display(),
        ])
    return rows


def export_csv(registrations, filename='report'):
    """Export registrations as CSV."""
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{filename}_{datetime.now():%Y%m%d}.csv"'

    writer = csv.writer(response)
    writer.writerow(REPORT_COLUMNS)

    for row in get_report_rows(registrations):
        writer.writerow(row)

    return response


def export_excel(registrations, filename='report'):
    """Export registrations as Excel XLSX."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Report'

    # Header styling
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2C5F7C', end_color='2C5F7C', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    # Write header
    for col, header in enumerate(REPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write data
    for row_idx, row in enumerate(get_report_rows(registrations), 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except (TypeError, AttributeError):
                pass
        adjusted_width = min(max_length + 4, 40)
        ws.column_dimensions[col[0].column_letter].width = adjusted_width

    # Freeze header row
    ws.freeze_panes = 'A2'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}_{datetime.now():%Y%m%d}.xlsx"'
    wb.save(response)
    return response


def export_pdf(registrations, filename='report', title='Saraswati Sanman Samaroh Report'):
    """Export registrations as PDF."""
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}_{datetime.now():%Y%m%d}.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=14,
        alignment=1,  # Center
        spaceAfter=12,
    )
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 5 * mm))

    # Prepare table data — use shorter column headers for PDF
    pdf_headers = [
        'Sr.', 'Student ID', 'Student Name', 'Family ID',
        'Family Name', 'Village', 'Std.', 'Medium',
        '%/CGPA', 'Achievement', 'Award', 'Status',
    ]

    rows = get_report_rows(registrations)
    # Remove 'Responsible Person' column (index 5) for PDF to save space
    pdf_rows = []
    for row in rows:
        pdf_row = row[:5] + row[6:]  # Skip index 5
        pdf_rows.append(pdf_row)

    table_data = [pdf_headers] + pdf_rows

    if len(table_data) > 1:
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C5F7C')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F7FA')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)
    else:
        elements.append(Paragraph('No records found for the selected filters.', styles['Normal']))

    doc.build(elements)
    return response
